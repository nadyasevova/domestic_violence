from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.common.logging_config import init_logging

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE_FILENAME = PROJECT_ROOT / "data/raw/population/districts_population.xlsx"



def extract_population_by_district(
    year: int,
    connection: Connection,
) -> pd.DataFrame:

    workbook = load_workbook(
        SOURCE_FILENAME,
        read_only=True,
        data_only=True,
    )

    records = []
    try:
        sheet_name = str(year)

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"No population sheet found for year {year}."
            )

        sheet = workbook[sheet_name]

        for row in sheet.iter_rows(min_row=6, max_col=4):
            name_cell = row[0]

            if not isinstance(name_cell.value, str):
                continue

            district_name = name_cell.value.strip()

            if not getattr(name_cell.font, "bold", False):
                continue

            if district_name == "Общо за страната":
                continue

            counts = [cell.value for cell in row[1:4]]

            if any(
                not isinstance(value, int) or value < 0
                for value in counts
            ):
                raise ValueError(
                    f"Invalid population values: "
                    f"year={year}, district={district_name}."
                )

            records.append({
                "population_date": date(year, 12, 31),
                "district_name_raw": district_name,
                "population_total": counts[0],
                "population_men": counts[1],
                "population_women": counts[2],
                "source_file": SOURCE_FILENAME.name,
                "source_sheet": sheet_name,
                "source_row_number": name_cell.row,
            })

    finally:
        workbook.close()

    if len(records) != 28:
        raise ValueError(
            f"Expected 28 districts for {year}, "
            f"found {len(records)}."
        )

    return pd.DataFrame(records)